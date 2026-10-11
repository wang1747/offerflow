import base64
import io
import zipfile

import pytest
from docx import Document
from PIL import Image

from app.modules.resume.parser import (
    _docx_xml_text,
    extract_docx_photo,
    parse_docx_bytes,
    parse_resume_bytes,
)
from app.modules.resume.structure import parse_resume_structure


def _docx_with_xml_text(text: str) -> bytes:
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>"
        + text
        + "</w:t></w:r></w:p></w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
    return buffer.getvalue()


def test_docx_xml_fallback_extracts_text():
    data = _docx_with_xml_text("文本框里的技能")
    assert "文本框里的技能" in _docx_xml_text(data)


def test_docx_xml_text_skips_mc_fallback_duplicate():
    """WPS 会在同一个段落内重复写入 Choice/Fallback 内容，只能读一次。"""
    choice = '<mc:Choice Requires="wps"><w:r><w:t>重复内容</w:t></w:r></mc:Choice>'
    fallback = "<mc:Fallback><w:r><w:t>重复内容</w:t></w:r></mc:Fallback>"
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        "<w:body><mc:AlternateContent>"
        "<w:p>" + choice + fallback + "</w:p>"
        "</mc:AlternateContent></w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    text = _docx_xml_text(buffer.getvalue())
    assert text.count("重复内容") == 1


def test_docx_xml_text_preserves_soft_line_breaks_and_tabs():
    """WPS 常在一个段落里用 br/tab 排版，导出前不能把它们压成空格。"""
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>姓名</w:t><w:br/><w:t>教育经历</w:t>"
        "<w:tab/><w:t>示例大学</w:t></w:r></w:p></w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    text = _docx_xml_text(buffer.getvalue())
    assert text.splitlines() == ["姓名", "教育经历 示例大学"]


def test_docx_xml_text_flattens_nested_wps_textbox_paragraphs():
    """WPS 外层段落包着文本框段落时，要按内层段落逐行输出。"""
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>外层重复文本</w:t></w:r><w:txbxContent>"
        "<w:p><w:r><w:t>第一行</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>第二行</w:t></w:r></w:p>"
        "</w:txbxContent></w:p></w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document_xml)

    text = _docx_xml_text(buffer.getvalue())
    assert text.splitlines() == ["第一行", "第二行"]


def test_parse_docx_bytes_reads_paragraph_and_table():
    doc = Document()
    doc.add_paragraph("示例大学")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "技能"
    table.rows[0].cells[1].text = "Python"
    buffer = io.BytesIO()
    doc.save(buffer)

    text = parse_docx_bytes(buffer.getvalue())
    assert "示例大学" in text
    assert "Python" in text


def test_parse_resume_bytes_empty_docx_raises():
    doc = Document()
    buffer = io.BytesIO()
    doc.save(buffer)
    with pytest.raises(ValueError):
        parse_resume_bytes("empty.docx", buffer.getvalue())


def test_plain_structure_does_not_treat_contact_lines_as_name():
    text = (
        "电话：18733444491\n"
        "电子邮箱：18733444491@163.com\n"
        "GitHub：https://github.com/wang1747\n"
        "教育经历\n"
        "示例大学 计算机科学与技术"
    )
    parsed = parse_resume_structure(text)
    assert parsed.name == ""
    assert len(parsed.contact) == 3
    assert parsed.sections[0].title == "教育经历"


def test_plain_structure_finds_name_before_contact_lines():
    text = (
        "张三\n"
        "电话：18733444491\n"
        "电子邮箱：18733444491@163.com\n"
        "GitHub：https://github.com/wang1747\n"
        "教育经历\n"
        "示例大学"
    )
    parsed = parse_resume_structure(text)
    assert parsed.name == "张三"
    assert len(parsed.contact) == 3


def test_docx_photo_over_2mb_is_compressed_before_storage():
    image = Image.effect_noise((3200, 3200), 100).convert("RGB")
    raw = io.BytesIO()
    image.save(raw, format="JPEG", quality=100)
    assert len(raw.getvalue()) > 2 * 1024 * 1024

    doc = Document()
    doc.add_paragraph("简历头像")
    doc.add_picture(io.BytesIO(raw.getvalue()))
    buffer = io.BytesIO()
    doc.save(buffer)

    photo = extract_docx_photo(buffer.getvalue())

    assert photo is not None
    assert photo["mime"] == "image/jpeg"
    assert len(base64.b64decode(photo["data"])) <= 2 * 1024 * 1024
