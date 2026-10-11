import pytest
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image
from PyPDF2 import PdfReader

from app.models.resume import Resume
from app.modules.resume.export_service import _bullet_line
from app.modules.resume_generation.contracts import GenerationResult
from app.modules.resume_generation.routes import _append_version


def test_resume_model_creation(db_session):
    """测试简历模型创建"""
    resume = Resume(
        user_id="test_user",
        version=1,
        raw_text="Test resume content",
        source_file="test.pdf",
    )
    db_session.add(resume)
    db_session.commit()
    assert resume.id is not None
    assert resume.version == 1


def test_resume_list_empty(db_session):
    """测试空简历列表"""
    resumes = db_session.query(Resume).all()
    assert len(resumes) == 0


def test_resume_versions_api(client, sample_jd_text):
    """测试同一份简历的版本历史（基础简历 + 优化版本归为一组）"""
    first = client.post(
        "/api/v1/resume/upload",
        json={"raw_text": "张三，3年Python开发经验，熟悉Django", "source_file": "a.txt"},
    )
    assert first.json()["success"] is True
    resume_id = first.json()["data"]["id"]

    optimize = client.post(
        "/api/v1/resume/optimize",
        json={"resume_id": resume_id, "jd_text": sample_jd_text},
    )
    assert optimize.status_code == 200

    versions = client.get(f"/api/v1/resume/{resume_id}/versions")
    assert versions.status_code == 200
    data = versions.json()["data"]
    assert len(data) == 2
    assert data[0]["version"] == 2
    assert data[0]["parsed_json"]["kind"] == "optimized"


def test_resume_upload_rejects_empty_text(client):
    """测试空文本简历不能上传"""
    response = client.post(
        "/api/v1/resume/upload",
        json={"raw_text": "   ", "source_file": "empty.txt"},
    )
    assert response.status_code == 400


def test_resume_upload_docx(client):
    """测试 .docx 简历上传"""
    doc = Document()
    doc.add_paragraph("张三，5年Python开发经验，熟悉Django和FastAPI")
    buffer = BytesIO()
    doc.save(buffer)

    response = client.post(
        "/api/v1/resume/upload-file",
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["filename"] == "resume.docx"

    versions = client.get(f"/api/v1/resume/{data['id']}/versions")
    assert versions.status_code == 200
    assert "Django" in versions.json()["data"][0]["raw_text"]


def test_resume_upload_docx_empty_rejected(client):
    """测试内容为空的 docx 上传会被拒绝"""
    doc = Document()
    buffer = BytesIO()
    doc.save(buffer)

    response = client.post(
        "/api/v1/resume/upload-file",
        files={
            "file": (
                "empty.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 400


def test_resume_upload_docx_keeps_photo_after_optimize_and_export(
    client, sample_resume_text, sample_jd_text,
):
    """优化后的版本必须继承原 DOCX 头像，并在 Word/PDF 导出里真正出现。"""
    photo = BytesIO()
    Image.new("RGB", (120, 160), (30, 58, 138)).save(photo, format="PNG")

    doc = Document()
    doc.add_paragraph(sample_resume_text)
    doc.add_picture(BytesIO(photo.getvalue()), width=Cm(2.5))
    buffer = BytesIO()
    doc.save(buffer)

    upload = client.post(
        "/api/v1/resume/upload-file",
        files={
            "file": (
                "resume-with-photo.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 200
    resume_id = upload.json()["data"]["id"]

    versions = client.get(f"/api/v1/resume/{resume_id}/versions")
    assert versions.status_code == 200
    source_photo = versions.json()["data"][0]["parsed_json"]["source_photo"]
    assert source_photo["mime"] == "image/png"
    assert source_photo["data"]

    optimize = client.post(
        "/api/v1/resume/optimize",
        json={"resume_id": resume_id, "jd_text": sample_jd_text},
    )
    assert optimize.status_code == 200
    optimized = optimize.json()["data"]
    assert optimized["new_version"] is not None
    optimized_id = optimized["new_version"]["id"]

    optimized_versions = client.get(f"/api/v1/resume/{optimized_id}/versions")
    optimized_payload = optimized_versions.json()["data"][0]["parsed_json"]
    assert optimized_payload["source_photo"] == source_photo

    word = client.get(f"/api/v1/resume/{optimized_id}/export?format=word")
    assert word.status_code == 200
    exported_doc = Document(BytesIO(word.content))
    assert len(exported_doc.inline_shapes) >= 1

    pdf = client.get(f"/api/v1/resume/{optimized_id}/export?format=pdf")
    assert pdf.status_code == 200
    page = PdfReader(BytesIO(pdf.content)).pages[0]
    xobjects = (page.get("/Resources") or {}).get("/XObject") or {}
    assert any(ref.get_object().get("/Subtype") == "/Image" for ref in xobjects.values())


def test_resume_export_pdf_and_word(client):
    """测试简历导出 PDF / Word"""
    upload = client.post(
        "/api/v1/resume/upload",
        json={
            "raw_text": "John Doe\nEducation: CS Master\nProject: Payment System\nSkills: Python, Docker",
            "source_file": "a.txt",
        },
    )
    resume_id = upload.json()["data"]["id"]

    pdf = client.get(f"/api/v1/resume/{resume_id}/export?format=pdf")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert len(pdf.content) > 100

    word = client.get(f"/api/v1/resume/{resume_id}/export?format=word")
    assert word.status_code == 200
    assert "wordprocessingml" in word.headers["content-type"]
    assert len(word.content) > 100


def test_resume_optimize_then_export(client, sample_resume_text, sample_jd_text):
    """测试优化后生成新版本并可直接导出，且不丢失原文"""
    upload = client.post(
        "/api/v1/resume/upload",
        json={"raw_text": sample_resume_text, "source_file": "a.txt"},
    )
    resume_id = upload.json()["data"]["id"]

    optimize = client.post(
        "/api/v1/resume/optimize",
        json={"resume_id": resume_id, "jd_text": sample_jd_text},
    )
    assert optimize.status_code == 200
    data = optimize.json()["data"]
    assert data["optimized"].strip() == sample_resume_text
    assert data["preservation"]["passed"] is True
    assert data["new_version"] is not None

    new_version_id = data["new_version"]["id"]
    word = client.get(f"/api/v1/resume/{new_version_id}/export?format=word")
    assert word.status_code == 200
    assert "wordprocessingml" in word.headers["content-type"]
    assert len(word.content) > 100


def test_optimized_word_keeps_source_photo_and_fills_a4_page(client, sample_jd_text):
    """上传 DOCX 的证件照要跨优化版本继承；Word 导出还要垂直铺满一页。"""
    photo = Image.new("RGB", (120, 160), "white")
    photo_buffer = BytesIO()
    photo.save(photo_buffer, format="PNG")

    source = Document()
    source.add_paragraph("张三，Python 后端开发，熟悉 FastAPI 和 PostgreSQL")
    source.add_picture(BytesIO(photo_buffer.getvalue()))
    source_buffer = BytesIO()
    source.save(source_buffer)

    upload = client.post(
        "/api/v1/resume/upload-file",
        files={
            "file": (
                "resume.docx",
                source_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 200, upload.text
    resume_id = upload.json()["data"]["id"]

    optimize = client.post(
        "/api/v1/resume/optimize",
        json={"resume_id": resume_id, "jd_text": sample_jd_text},
    )
    assert optimize.status_code == 200, optimize.text
    new_version_id = optimize.json()["data"]["new_version"]["id"]

    versions = client.get(f"/api/v1/resume/{resume_id}/versions").json()["data"]
    latest = next(v for v in versions if v["id"] == new_version_id)
    assert latest["parsed_json"]["source_photo"]["data"]

    exported = client.get(f"/api/v1/resume/{new_version_id}/export?format=word")
    assert exported.status_code == 200, exported.text
    document = Document(BytesIO(exported.content))
    assert len(document.inline_shapes) == 1
    assert round(document.sections[0].page_width.cm, 1) == 21.0
    vertical = document.sections[0]._sectPr.find(qn("w:vAlign"))
    assert vertical is not None
    assert vertical.get(qn("w:val")) == "both"


def test_delete_resume_removes_every_version(
    client, sample_resume_text, sample_jd_text,
):
    upload = client.post(
        "/api/v1/resume/upload",
        json={"raw_text": sample_resume_text, "source_file": "a.txt"},
    )
    resume_id = upload.json()["data"]["id"]
    optimize = client.post(
        "/api/v1/resume/optimize",
        json={"resume_id": resume_id, "jd_text": sample_jd_text},
    )
    latest_id = optimize.json()["data"]["new_version"]["id"]

    deleted = client.delete(f"/api/v1/resume/{latest_id}")

    assert deleted.status_code == 200
    assert deleted.json()["data"]["deleted"] == 2
    assert client.get("/api/v1/resume/list").json()["data"] == []


def test_append_generation_version_inherits_source_photo(db_session):
    photo = {"data": "cGhvdG8=", "mime": "image/png"}
    source = Resume(
        user_id="test_user",
        version=1,
        raw_text="张三，Python 后端开发",
        parsed_json={"source_photo": photo},
        source_file="resume.docx",
    )
    db_session.add(source)
    db_session.commit()

    new_version = _append_version(
        db_session,
        "test_user",
        source,
        GenerationResult(
            name="张三",
            position="Python 后端开发",
            resume_text="张三，Python 后端开发",
        ),
    )

    assert new_version.parsed_json["source_photo"] == photo


def test_export_cleans_triangle_bullet_marks():
    assert _bullet_line("▸ Python 后端开发") == "Python 后端开发"
    assert _bullet_line("▶ FastAPI") == "FastAPI"
    assert _bullet_line("► Docker") == "Docker"
