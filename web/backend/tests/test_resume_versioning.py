"""简历版本管理：一份简历 = 一个 document，每次变更只追加不可变版本。

守的三条线：
1. 新上传是**独立文档**（v1），不再按用户全局递增；
2. 优化 / 逐段重写 / 候选保存 / 回滚**只追加**，旧版本一个字节都不改；
3. 导出补齐结构化数据**不写库**（导出是只读动作，写库等于偷偷改历史版本）。
"""

import base64

from docx import Document
from io import BytesIO
from PIL import Image

from test_resume_generation import _generate_payload

from app.core.database import get_db
from app.modules.resume.versioning import document_id
from app.models.resume import Resume


def _upload(client, text: str) -> dict:
    response = client.post(
        "/api/v1/resume/upload",
        json={"raw_text": text, "source_file": "a.txt"},
    )
    assert response.status_code == 200
    return response.json()["data"]


def _versions(client, resume_id: str) -> list:
    response = client.get(f"/api/v1/resume/{resume_id}/versions")
    assert response.status_code == 200
    return response.json()["data"]


def _snapshot(client, resume_id: str) -> dict:
    """按 id 记录每个版本的 raw_text / parsed_json，用于证明"旧版本没被动过"。"""
    return {v["id"]: (v["raw_text"], v["parsed_json"]) for v in _versions(client, resume_id)}


def _generate(client) -> dict:
    response = client.post("/api/v1/resume-generation/generate", json=_generate_payload())
    assert response.json()["success"] is True
    return response.json()["data"]


def test_new_upload_is_its_own_document_starting_at_v1(client):
    a = _upload(client, "张三，3年 Python 开发经验，熟悉 Django")
    b = _upload(client, "李四，2年 Java 开发经验，熟悉 Spring")

    # 不再按用户全局递增：两份互不相关的简历都从 v1 开始
    assert a["version"] == 1
    assert b["version"] == 1

    versions_a = _versions(client, a["id"])
    versions_b = _versions(client, b["id"])
    assert len(versions_a) == 1 and len(versions_b) == 1
    assert versions_a[0]["parsed_json"]["document_id"] == a["id"]
    assert versions_b[0]["parsed_json"]["document_id"] == b["id"]
    assert versions_a[0]["parsed_json"]["document_id"] != versions_b[0]["parsed_json"]["document_id"]


def test_content_change_appends_version_and_never_touches_the_old_one(client):
    generated = _generate(client)
    before = _snapshot(client, generated["id"])

    response = client.post(
        "/api/v1/resume-generation/regenerate-section",
        json={
            "resume_id": generated["id"],
            "structured": {
                "name": generated["name"],
                "position": generated["position"],
                "contact": generated["contact"],
                "summary": generated["summary"],
                "education": generated["education"],
                "experiences": generated["experiences"],
                "skills": generated["skills"],
                "certifications": generated["certifications"],
            },
            "section": "summary",
            "jd_text": "Python 后端实习，要求熟悉 FastAPI",
        },
    )
    assert response.json()["success"] is True
    new_data = response.json()["data"]

    # 追加而非原地改写
    assert new_data["id"] != generated["id"]
    assert new_data["version"] == generated["version"] + 1
    assert new_data["document_id"] == generated["id"]

    after = _snapshot(client, generated["id"])
    assert len(after) == len(before) + 1
    for old_id, snapshot in before.items():
        assert after[old_id] == snapshot, "旧版本被改写了 —— 版本必须不可变"


def test_optimize_appends_version_and_list_groups_by_document(client, sample_jd_text):
    uploaded = _upload(client, "张三，3年 Python 开发经验，熟悉 Django")

    optimize = client.post(
        "/api/v1/resume/optimize",
        json={"resume_id": uploaded["id"], "jd_text": sample_jd_text},
    )
    assert optimize.status_code == 200
    assert optimize.json()["data"]["new_version"] is not None

    versions = _versions(client, uploaded["id"])
    assert [v["version"] for v in versions] == [2, 1]
    assert versions[0]["parsed_json"]["document_id"] == uploaded["id"]

    # 列表按 document_id 聚合：同一份简历只出现一条，并带上版本数
    items = client.get("/api/v1/resume/list").json()["data"]
    assert len(items) == 1
    assert items[0]["version_count"] == 2
    assert items[0]["document_id"] == uploaded["id"]
    assert items[0]["version"] == 2


def test_rollback_appends_new_version_and_records_source(client, sample_jd_text):
    uploaded = _upload(client, "张三，3年 Python 开发经验，熟悉 Django")
    client.post("/api/v1/resume/optimize", json={"resume_id": uploaded["id"], "jd_text": sample_jd_text})
    before = _snapshot(client, uploaded["id"])

    # 回滚到 v1
    versions = _versions(client, uploaded["id"])
    first_version_id = [v for v in versions if v["version"] == 1][0]["id"]

    response = client.post(f"/api/v1/resume/{uploaded['id']}/versions/{first_version_id}/rollback")
    assert response.status_code == 200
    rolled = response.json()["data"]
    assert rolled["version"] == 3  # 追加一条新的最高版本，而不是把 v1 变成当前

    after = _snapshot(client, uploaded["id"])
    assert len(after) == len(before) + 1
    for old_id, snapshot in before.items():
        assert after[old_id] == snapshot, "回滚不许改动历史版本"

    new_row = [v for v in _versions(client, uploaded["id"]) if v["id"] == rolled["id"]][0]
    assert new_row["parsed_json"]["rolled_back_from"] == 1
    assert new_row["parsed_json"]["document_id"] == uploaded["id"]


def test_cross_document_rollback_is_rejected(client):
    a = _upload(client, "张三，3年 Python 开发经验")
    b = _upload(client, "李四，2年 Java 开发经验")

    response = client.post(f"/api/v1/resume/{a['id']}/versions/{b['id']}/rollback")
    assert response.status_code == 400
    assert "同一份简历" in response.json()["detail"]
    # 拒绝之后两份文档各自还是 1 版，没有被污染
    assert len(_versions(client, a["id"])) == 1
    assert len(_versions(client, b["id"])) == 1


def test_export_fills_structured_in_memory_without_writing_back(client, _fake_llm):
    """导出补齐结构化数据：既要用上（导出层次正确），又绝不能写库（版本不可变）。"""
    calls = {"n": 0}

    def _extract(text: str) -> dict:
        calls["n"] += 1
        return {
            "name": "提取姓名",
            "position": "提取意向",
            "contact": {"phone": "", "email": "", "github": ""},
            "summary": "提取出来的总结",
            "education": [],
            "experiences": [
                {"type": "project", "name": "提取项目", "role": "开发", "start": "", "end": "", "bullets": ["提取出来的要点"]}
            ],
            "skills": ["Python"],
            "certifications": [],
        }

    _fake_llm.extract_handler = _extract
    uploaded = _upload(client, "John Doe\nEducation: CS Master\nSkills: Python, Docker")
    before = _snapshot(client, uploaded["id"])

    exported = client.get(f"/api/v1/resume/{uploaded['id']}/export?format=word")
    assert exported.status_code == 200
    assert len(exported.content) > 100

    # ① 真的用上了内存里的结构化（否则这次导出会退化成纯文本）
    text = "\n".join(p.text for p in Document(BytesIO(exported.content)).paragraphs)
    assert "提取姓名" in text

    # ② 一个字节都没写回数据库
    assert _snapshot(client, uploaded["id"]) == before
    stored = _versions(client, uploaded["id"])[0]["parsed_json"]
    assert "structured" not in stored and "generation" not in stored

    # ③ 同一份简历第二次导出直接命中进程内缓存，不再花第二次模型调用
    again = client.get(f"/api/v1/resume/{uploaded['id']}/export?format=pdf")
    assert again.status_code == 200
    assert calls["n"] == 1


def test_document_id_helper_follows_parent_chain_for_legacy_rows(db_session):
    """老数据没有 document_id 时，沿 parent_id 回溯到根 —— 保证兼容、不需要迁移。"""
    from app.models.resume import Resume

    root = Resume(
        id="legacy-root",
        user_id="default",
        version=1,
        raw_text="legacy",
        parsed_json={"sections": []},
    )
    child = Resume(
        id="legacy-child",
        user_id="default",
        version=2,
        raw_text="legacy optimized",
        parsed_json={"parent_id": "legacy-root", "kind": "optimized"},
    )
    grandchild = Resume(
        id="legacy-grandchild",
        user_id="default",
        version=3,
        raw_text="legacy again",
        parsed_json={"parent_id": "legacy-child", "kind": "optimized"},
    )
    db_session.add_all([root, child, grandchild])
    db_session.commit()

    mapping = {row.id: row for row in (root, child, grandchild)}
    assert document_id(grandchild, mapping) == "legacy-root"
    assert document_id(child, mapping) == "legacy-root"
    assert document_id(root, mapping) == "legacy-root"


def test_export_word_recovers_photo_from_parent_for_early_optimized_version(client):
    """早期优化版本漏存 source_photo 时，导出要能从同文档根版本原地恢复。"""
    image = Image.new("RGB", (120, 160), "white")
    image_buffer = BytesIO()
    image.save(image_buffer, format="PNG")
    photo = {
        "data": base64.b64encode(image_buffer.getvalue()).decode("ascii"),
        "mime": "image/png",
    }
    root = Resume(
        id="legacy-photo-root",
        user_id="default",
        version=1,
        raw_text="张三，Python 后端开发",
        parsed_json={"document_id": "legacy-photo-root", "source_photo": photo},
        source_file="resume.docx",
    )
    child = Resume(
        id="legacy-photo-child",
        user_id="default",
        version=2,
        raw_text="张三，Python 后端开发",
        parsed_json={
            "document_id": "legacy-photo-root",
            "parent_id": "legacy-photo-root",
            "kind": "optimized",
            "structured": {
                "name": "张三",
                "position": "Python 后端开发",
                "contact": {},
                "summary": "",
                "education": [],
                "experiences": [],
                "skills": [],
                "certifications": [],
            },
        },
        source_file="resume.docx",
    )
    session = next(client.app.dependency_overrides[get_db]())
    try:
        session.add_all([root, child])
        session.commit()
    finally:
        session.close()

    exported = client.get("/api/v1/resume/legacy-photo-child/export?format=word")

    assert exported.status_code == 200
    assert len(Document(BytesIO(exported.content)).inline_shapes) == 1
