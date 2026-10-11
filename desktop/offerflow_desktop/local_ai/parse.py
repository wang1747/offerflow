"""本地文档解析（P1）。

分层策略（2026-09-30 真机实测后定稿，见 offerflow_desktop_plan.md §5 P1）：

    PDF    → LiteParse（Apache-2.0，自带 PDFium，免模型，实测 ~13ms/页，markdown 结构优于 pypdf）
             ├─ 抽不出文字（扫描件）→ LiteParse.screenshot() 渲染页图 → Windows 系统 OCR
             └─ LiteParse 异常 → pypdf 兜底
    DOCX   → python-docx（段落 + 表格）
    TXT/MD → 直读（utf-8 → gbk → utf-16 探测）
    图片   → Windows 系统 OCR（core/ocr.py）

为什么**不用 PyMuPDF4LLM**：它继承 PyMuPDF 的 AGPL-3.0 传染性许可，闭源分发有合规风险
（上游博客亦点名此点）；且 PDF 一路已被 LiteParse 覆盖——同基准下 LiteParse 更快更准。

为什么**不用 MarkItDown**：会连带安装 onnxruntime + magika + sympy 等（实测约 60MB+，
其中 magika 还带 AI 文件类型识别模型），与「轻量、体积可控」目标冲突；而它主打的
Office 格式里，DOCX 已由 python-docx 覆盖（XLSX/PPTX 桌面版当前用不到）。

为什么 **DOCX 不走 LiteParse**：实测 LiteParse 处理 Office 文件是转调 LibreOffice
（`ParseError: LibreOffice is not installed`），需额外装几百 MB 桌面软件，不划算。

为什么**扫描件用 Windows 系统 OCR**：liteparse 的 `ocr_enabled=True` 走 Tesseract
（需 tessdata 二进制），而本项目已有免模型的 winrt 系统 OCR（core/ocr.py，P1 已加固）。

打包注意（PyInstaller）：liteparse 带原生文件 `pdfium.dll`(7MB) 与 `_liteparse.pyd`(20MB)，
**只写 hiddenimports 会漏**，须 `collect_all("liteparse")`（见 OfferFlow.spec）。
"""
from __future__ import annotations

import os
import time
from dataclasses import dataclass
from typing import Dict, List, Optional

from ..core.ocr import is_ocr_available, ocr_image_file

# 扫描件判定：LiteParse 抽出的有效字符数低于「每页 15 字」时，认为该 PDF 没有文本层
_SCAN_CHARS_PER_PAGE = 15

# 支持的文件类型（UI 文件选择器用）
SUPPORTED_EXT = (".pdf", ".docx", ".txt", ".md", ".markdown",
                 ".png", ".jpg", ".jpeg", ".bmp", ".webp")
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".bmp", ".webp")


@dataclass
class ParseResult:
    """统一的解析结果。"""

    text: str = ""
    engine: str = ""          # liteparse / pypdf / ocr / python-docx / plain
    pages: int = 0
    elapsed_ms: int = 0
    degraded: bool = False    # 是否走了非首选路径
    note: str = ""            # 给用户/日志看的一句话说明

    @property
    def empty(self) -> bool:
        return not self.text.strip()


# --------------------------------------------------------------------------- #
# 引擎可用性
# --------------------------------------------------------------------------- #
def _liteparse_available() -> bool:
    try:
        import liteparse  # noqa: F401
        return True
    except Exception:  # noqa: BLE001
        return False


def _pypdf_available() -> bool:
    try:
        import pypdf  # noqa: F401
        return True
    except Exception:  # noqa: BLE001
        return False


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except Exception:  # noqa: BLE001
        return False


def parse_status() -> Dict[str, object]:
    """各解析引擎可用性（供 UI / 自检显示，等价于「装库≠生效」的实测字段）。"""
    return {
        "liteparse": _liteparse_available(),
        "pypdf": _pypdf_available(),
        "python-docx": _docx_available(),
        "windows-ocr": is_ocr_available(),
    }


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _pdf_via_liteparse(path: str) -> ParseResult:
    from liteparse import LiteParse

    t0 = time.time()
    # ocr_enabled=False：LiteParse 自带 OCR 走 Tesseract，缺失时会**联网去 GitHub 下载
    # tessdata**（实测日志 `error sending request for url .../tessdata_best/eng.traineddata`）。
    # 本项目 OCR 统一走 Windows 系统引擎（core/ocr.py），关掉以免产生意外联网与超时。
    res = LiteParse(output_format="markdown", quiet=True, ocr_enabled=False).parse(path)
    text = (res.text or "").strip()
    pages = len(getattr(res, "pages", None) or [])
    return ParseResult(
        text=text,
        engine="liteparse",
        pages=pages,
        elapsed_ms=int((time.time() - t0) * 1000),
    )


def _pdf_via_pypdf(path: str) -> ParseResult:
    from pypdf import PdfReader

    t0 = time.time()
    reader = PdfReader(path)
    text = "\n".join(
        (page.extract_text() or "").strip() for page in reader.pages
    ).strip()
    return ParseResult(
        text=text,
        engine="pypdf",
        pages=len(reader.pages),
        elapsed_ms=int((time.time() - t0) * 1000),
        degraded=True,
        note="LiteParse 不可用或失败，已降级到 pypdf（结构信息较少）",
    )


def _pdf_via_ocr(path: str) -> ParseResult:
    """扫描件 PDF：逐页渲染成位图后交给 Windows 系统 OCR。

    渲染依赖 LiteParse.screenshot()（它自带 PDFium）；若 LiteParse 不可用则无法渲染，
    返回空结果并说明原因。
    """
    if not is_ocr_available():
        return ParseResult(engine="ocr", note="本机未启用 Windows OCR 引擎")

    t0 = time.time()
    try:
        from liteparse import LiteParse

        # dpi=300：扫描件渲染分辨率。真实扫描件多为 200~300dpi 位图，300dpi 渲染即 1:1、
        # 不丢像素；但不能指望「dpi 越高越准」——实测把一张小尺寸 UI 截图放大到 300dpi 后
        # 因上采样插值反而更糊（识别字数 66 → 45）。OCR 精度取决于**有效像素**，非 dpi 数值。
        shots = LiteParse(quiet=True, ocr_enabled=False, dpi=300).screenshot(path)
    except Exception as exc:  # noqa: BLE001
        return ParseResult(engine="ocr", note=f"页面渲染失败：{type(exc).__name__}")

    import tempfile
    import uuid

    chunks: List[str] = []
    for shot in shots:
        tmp = os.path.join(tempfile.gettempdir(), f"offerflow_scan_{uuid.uuid4().hex}.png")
        try:
            with open(tmp, "wb") as fh:
                fh.write(shot.image_bytes)
            page_text = ocr_image_file(tmp)
        except Exception:  # noqa: BLE001
            page_text = ""
        finally:
            try:
                os.remove(tmp)
            except Exception:  # noqa: BLE001
                pass
        if page_text.strip():
            chunks.append(page_text.strip())

    text = "\n\n".join(chunks).strip()
    return ParseResult(
        text=text,
        engine="ocr",
        pages=len(shots),
        elapsed_ms=int((time.time() - t0) * 1000),
        degraded=True,
        note="该 PDF 无文本层，已用系统 OCR 逐页识别（可能有错字，请核对）",
    )


def parse_pdf(path: str) -> ParseResult:
    """PDF：LiteParse → 扫描件 OCR → pypdf 兜底。"""
    result: Optional[ParseResult] = None

    if _liteparse_available():
        try:
            result = _pdf_via_liteparse(path)
        except Exception as exc:  # noqa: BLE001
            result = ParseResult(
                engine="liteparse",
                degraded=True,
                note=f"LiteParse 失败：{type(exc).__name__}: {exc}",
            )

    # 无文本层 → 扫描件，尝试 OCR
    pages = result.pages if result else 0
    if (result is None or len(result.text) < _SCAN_CHARS_PER_PAGE * max(pages, 1)):
        ocr_result = _pdf_via_ocr(path)
        if len(ocr_result.text) > len(result.text if result else ""):
            # 保留 LiteParse 的页数信息
            if result and pages:
                ocr_result.pages = ocr_result.pages or pages
            return ocr_result
        if result and result.text.strip():
            return result

    if result and result.text.strip():
        return result

    # 最后兜底：pypdf
    if _pypdf_available():
        try:
            return _pdf_via_pypdf(path)
        except Exception as exc:  # noqa: BLE001
            return ParseResult(
                engine="pypdf",
                degraded=True,
                note=f"pypdf 读取失败：{type(exc).__name__}",
            )

    return result or ParseResult(engine="", note="无可用的 PDF 解析引擎")


# --------------------------------------------------------------------------- #
# DOCX / 纯文本 / 图片
# --------------------------------------------------------------------------- #
def parse_docx(path: str) -> ParseResult:
    """Word：段落 + 表格（表格按行拼成 `a | b | c`）。"""
    t0 = time.time()
    import docx

    document = docx.Document(path)
    parts: List[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return ParseResult(
        text="\n".join(parts).strip(),
        engine="python-docx",
        elapsed_ms=int((time.time() - t0) * 1000),
    )


def parse_plain(path: str) -> ParseResult:
    """纯文本：按 utf-8 → gbk → utf-16 顺序探测编码。"""
    t0 = time.time()
    text = ""
    for enc in ("utf-8", "gbk", "utf-16"):
        try:
            with open(path, encoding=enc) as fh:
                text = fh.read()
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if not text:
        with open(path, encoding="utf-8", errors="ignore") as fh:
            text = fh.read()
    return ParseResult(
        text=text.strip(),
        engine="plain",
        elapsed_ms=int((time.time() - t0) * 1000),
    )


def parse_image(path: str) -> ParseResult:
    """图片：Windows 系统 OCR。"""
    t0 = time.time()
    if not is_ocr_available():
        return ParseResult(engine="ocr", note="本机未启用 Windows OCR 引擎")
    text = ocr_image_file(path)
    return ParseResult(
        text=text.strip(),
        engine="ocr",
        degraded=True,
        elapsed_ms=int((time.time() - t0) * 1000),
        note="图片经系统 OCR 识别（可能有错字，请核对）",
    )


# --------------------------------------------------------------------------- #
# 主入口
# --------------------------------------------------------------------------- #
def parse_document(path: str) -> ParseResult:
    """按扩展名分派到对应解析器。文件不存在 / 类型不支持时抛 ValueError。"""
    if not path or not os.path.isfile(path):
        raise ValueError(f"文件不存在：{path}")

    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        if not _docx_available():
            raise ValueError("未安装 python-docx，无法解析 Word 文件")
        return parse_docx(path)
    if ext in (".txt", ".md", ".markdown"):
        return parse_plain(path)
    if ext in IMAGE_EXT:
        return parse_image(path)
    raise ValueError(f"不支持的文件类型：{ext or '未知'}")


def extract_text(path: str) -> str:
    """便捷封装：只要文本，失败 / 空结果抛 ValueError（UI 直接捕获展示）。"""
    result = parse_document(path)
    if result.empty:
        if result.note:
            raise ValueError(result.note)
        raise ValueError("未能从该文件提取到文本（可能是扫描版且本机 OCR 不可用）")
    return result.text
